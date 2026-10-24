import openpyxl
import os

files = [
    r'C:\Users\XIAO\Desktop\大类信息\2025河北单招_分数线.xlsx',
    r'C:\Users\XIAO\Desktop\大类信息\2025河北单招_招生计划_含学费.xlsx',
    r'C:\Users\XIAO\Desktop\大类信息\2025河北单招_一分一档.xlsx',
]

for f in files:
    print(f'\n=== {os.path.basename(f)} ===')
    print(f'File size: {os.path.getsize(f)} bytes')
    wb = openpyxl.load_workbook(f)
    print(f'Sheets: {wb.sheetnames}')
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        print(f'  {sheet_name}: {ws.max_row} rows x {ws.max_column} cols')
        # Print first few rows
        if sheet_name != '说明':
            for row in ws.iter_rows(min_row=5, max_row=7, values_only=True):
                print(f'    {row}')
    wb.close()
