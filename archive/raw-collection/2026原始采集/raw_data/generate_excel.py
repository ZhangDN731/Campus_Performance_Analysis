#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成4个Excel工作簿：分数线、招生计划、一分一档、院校详情"""

import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

OUT_DIR = r'C:\Users\XIAO\Desktop\大类信息'
RAW_DIR = os.path.join(OUT_DIR, 'raw_data')

# 样式定义
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
CELL_FONT = Font(name='微软雅黑', size=10)
CELL_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)
THIN_BORDER = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
TITLE_FONT = Font(name='微软雅黑', size=14, bold=True)
NOTE_FONT = Font(name='微软雅黑', size=10)

def get_sheet_name(cat_name):
    """将大类名称转换为合法的sheet名称（不超过31字符）"""
    mapping = {
        '考试一类': '一类',
        '考试二类': '二类',
        '考试三类': '三类',
        '考试四类': '四类',
        '考试五类': '五类',
        '考试六类': '六类',
        '考试七类': '七类',
        '考试八类-美术设计类': '八类美术',
        '考试九类': '九类',
        '考试十类': '十类',
    }
    return mapping.get(cat_name, cat_name[:31])

def style_header_row(ws, row_num, num_cols):
    """设置表头行样式"""
    for col in range(1, num_cols + 1):
        cell = ws.cell(row=row_num, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER

def style_data_cells(ws, start_row, end_row, num_cols, left_cols=None):
    """设置数据单元格样式"""
    if left_cols is None:
        left_cols = []
    for row in range(start_row, end_row + 1):
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.font = CELL_FONT
            cell.border = THIN_BORDER
            if col in left_cols:
                cell.alignment = LEFT_ALIGN
            else:
                cell.alignment = CELL_ALIGN

def auto_column_width(ws, num_cols, min_width=10, max_width=50):
    """自动调整列宽"""
    for col in range(1, num_cols + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col, values_only=True):
            cell_value = str(row[0]) if row[0] else ''
            # 中文字符算2个宽度
            length = sum(2 if ord(c) > 127 else 1 for c in cell_value)
            max_len = max(max_len, length)
        width = min(max(max_len + 4, min_width), max_width)
        ws.column_dimensions[get_column_letter(col)].width = width

def create_info_sheet(wb, title, data_source, control_lines=None, notes=None):
    """创建说明sheet"""
    ws = wb.create_sheet('说明', 0)
    ws.sheet_properties.tabColor = 'FFC000'
    
    row = 1
    ws.cell(row=row, column=1, value=title).font = TITLE_FONT
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    row += 2
    
    ws.cell(row=row, column=1, value='数据来源').font = Font(name='微软雅黑', size=10, bold=True)
    ws.cell(row=row, column=2, value=data_source).font = NOTE_FONT
    row += 1
    
    ws.cell(row=row, column=1, value='采集时间').font = Font(name='微软雅黑', size=10, bold=True)
    ws.cell(row=row, column=2, value='2026-10-20').font = NOTE_FONT
    row += 1
    
    ws.cell(row=row, column=1, value='数据年份').font = Font(name='微软雅黑', size=10, bold=True)
    ws.cell(row=row, column=2, value='2025年').font = NOTE_FONT
    row += 2
    
    if control_lines:
        ws.cell(row=row, column=1, value='各类省控制线汇总').font = Font(name='微软雅黑', size=11, bold=True)
        row += 1
        ws.cell(row=row, column=1, value='考试类').font = HEADER_FONT
        ws.cell(row=row, column=1).fill = HEADER_FILL
        ws.cell(row=row, column=2, value='省控制线').font = HEADER_FONT
        ws.cell(row=row, column=2).fill = HEADER_FILL
        ws.cell(row=row, column=3, value='涉及专业大类').font = HEADER_FONT
        ws.cell(row=row, column=3).fill = HEADER_FILL
        for col in range(1, 4):
            ws.cell(row=row, column=col).alignment = HEADER_ALIGN
            ws.cell(row=row, column=col).border = THIN_BORDER
        row += 1
        for cat_name, line, desc in control_lines:
            ws.cell(row=row, column=1, value=cat_name).font = CELL_FONT
            ws.cell(row=row, column=2, value=line).font = CELL_FONT
            ws.cell(row=row, column=3, value=desc).font = CELL_FONT
            for col in range(1, 4):
                ws.cell(row=row, column=col).border = THIN_BORDER
                ws.cell(row=row, column=col).alignment = CELL_ALIGN
            row += 1
        row += 1
    
    if notes:
        ws.cell(row=row, column=1, value='数据说明').font = Font(name='微软雅黑', size=11, bold=True)
        row += 1
        for note in notes:
            ws.cell(row=row, column=1, value=note).font = NOTE_FONT
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
            row += 1
    
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 15
    ws.column_dimensions['C'].width = 50
    ws.column_dimensions['D'].width = 15

def generate_score_lines():
    """生成分数线工作簿"""
    print('=== 生成分数线工作簿 ===')
    with open(os.path.join(RAW_DIR, 'score_lines.json'), 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    wb = Workbook()
    wb.remove(wb.active)
    
    control_lines = []
    for cid, cat in data.items():
        control_lines.append((cat['name'], cat['control_line'], cat['desc']))
    
    notes = [
        '1. 本表数据来源于河北单招网(www.hbdzxx.com)，为2025年河北省普通高职单招各考试类录取投档分数线。',
        '2. "录取投档分"列全部显示为"—"，表示河北省教育考试院尚未公布各院校专业的实际投档分数线。',
        '3. 省控制线为各考试类的最低录取控制分数线，考生成绩达到此线方可参与投档。',
        '4. 专业名称中括号内为校区、就读地点、身体要求等备注信息，请仔细阅读。',
        '5. 数据仅供参考，最终以河北省教育考试院及各院校官方公布为准。',
    ]
    
    create_info_sheet(wb, '2025年河北单招各类录取分数线', 'http://www.hbdzxx.com/score/line/2025/', control_lines, notes)
    
    for cid, cat in data.items():
        sheet_name = get_sheet_name(cat['name'])
        ws = wb.create_sheet(sheet_name)
        
        # 顶部标注控制线
        ws.cell(row=1, column=1, value=f'{cat["name"]}（{cat["desc"]}）').font = TITLE_FONT
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        ws.cell(row=2, column=1, value=f'省控制线：{cat["control_line"]}分').font = Font(name='微软雅黑', size=11, bold=True, color='FF0000')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
        ws.cell(row=3, column=1, value=f'数据来源：{cat["source_url"]}').font = NOTE_FONT
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=5)
        
        # 表头
        headers = ['序号', '院校名称', '专业名称', '考试类', '录取投档分']
        for col, h in enumerate(headers, 1):
            ws.cell(row=5, column=col, value=h)
        style_header_row(ws, 5, len(headers))
        
        # 数据
        for i, row_data in enumerate(cat['data'], 1):
            ws.cell(row=5 + i, column=1, value=i)
            ws.cell(row=5 + i, column=2, value=row_data[0])
            ws.cell(row=5 + i, column=3, value=row_data[1])
            ws.cell(row=5 + i, column=4, value=row_data[2])
            ws.cell(row=5 + i, column=5, value=row_data[3])
        
        style_data_cells(ws, 6, 5 + len(cat['data']), len(headers), left_cols=[2, 3])
        auto_column_width(ws, len(headers))
        ws.freeze_panes = 'A6'
        
        print(f'  {sheet_name}: {len(cat["data"])} 条')
    
    filepath = os.path.join(OUT_DIR, '2025河北单招_分数线.xlsx')
    wb.save(filepath)
    print(f'  保存: {filepath}')
    return filepath

def generate_plans():
    """生成招生计划工作簿"""
    print('=== 生成招生计划工作簿 ===')
    with open(os.path.join(RAW_DIR, 'plans.json'), 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    wb = Workbook()
    wb.remove(wb.active)
    
    control_lines = []
    total_plans = 0
    for cid, cat in data.items():
        control_lines.append((cat['name'], cat['control_line'], cat['desc']))
        for row in cat['data']:
            try:
                total_plans += int(row[3])
            except:
                pass
    
    notes = [
        f'1. 本表数据来源于河北单招网(www.hbdzxx.com)，为2025年河北省普通高职单招各考试类招生计划。',
        f'2. 十大类招生计划总计约 {total_plans} 人（含中外合作办学等特殊类型）。',
        '3. "学费(元/年)"为院校公布的学费标准，部分专业学费可能为暂定，最终以院校官方公布为准。',
        '4. 专业名称中括号内为校区、就读地点、身体要求、校企合作、中外合作办学等备注信息。',
        '5. 学制均为3年（部分专业可能有差异，以院校公布为准）。',
        '6. 数据仅供参考，最终以河北省教育考试院及各院校官方公布为准。',
    ]
    
    create_info_sheet(wb, '2025年河北单招各类招生计划（含专业计划及学费）', 'http://www.hbdzxx.com/school/plan/2025/', control_lines, notes)
    
    for cid, cat in data.items():
        sheet_name = get_sheet_name(cat['name'])
        ws = wb.create_sheet(sheet_name)
        
        # 顶部标注
        ws.cell(row=1, column=1, value=f'{cat["name"]}（{cat["desc"]}）').font = TITLE_FONT
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=7)
        ws.cell(row=2, column=1, value=f'省控制线：{cat["control_line"]}分').font = Font(name='微软雅黑', size=11, bold=True, color='FF0000')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=7)
        ws.cell(row=3, column=1, value=f'数据来源：{cat["source_url"]}').font = NOTE_FONT
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=7)
        
        # 表头
        headers = ['序号', '院校名称', '专业名称', '考试类', '计划数', '学制', '学费(元/年)']
        for col, h in enumerate(headers, 1):
            ws.cell(row=5, column=col, value=h)
        style_header_row(ws, 5, len(headers))
        
        # 数据
        for i, row_data in enumerate(cat['data'], 1):
            ws.cell(row=5 + i, column=1, value=i)
            ws.cell(row=5 + i, column=2, value=row_data[0])
            ws.cell(row=5 + i, column=3, value=row_data[1])
            ws.cell(row=5 + i, column=4, value=row_data[2])
            # 计划数转为数字
            try:
                ws.cell(row=5 + i, column=5, value=int(row_data[3]))
            except:
                ws.cell(row=5 + i, column=5, value=row_data[3])
            ws.cell(row=5 + i, column=6, value=row_data[4])
            # 学费转为数字
            try:
                ws.cell(row=5 + i, column=7, value=int(row_data[5]))
            except:
                ws.cell(row=5 + i, column=7, value=row_data[5])
        
        style_data_cells(ws, 6, 5 + len(cat['data']), len(headers), left_cols=[2, 3])
        auto_column_width(ws, len(headers))
        ws.freeze_panes = 'A6'
        
        print(f'  {sheet_name}: {len(cat["data"])} 条')
    
    filepath = os.path.join(OUT_DIR, '2025河北单招_招生计划_含学费.xlsx')
    wb.save(filepath)
    print(f'  保存: {filepath}')
    return filepath

def generate_yifenyidang():
    """生成一分一档工作簿"""
    print('=== 生成一分一档工作簿 ===')
    with open(os.path.join(RAW_DIR, 'yifenyidang.json'), 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    wb = Workbook()
    wb.remove(wb.active)
    
    notes = [
        '1. 本表数据来源于河北单招网(www.hbdzxx.com)，为2025年河北省普通高职单招各考试类一分一档统计表。',
        '2. "分数"列首行为"XXX及以上"，表示该分数及以上的考生人数汇总。',
        '3. "本段人数"为该分数对应的考生人数，"累计人数"为该分数及以上的考生总人数。',
        '4. "累计占比"为累计人数占该考试类总考生人数的百分比。',
        '5. 数据仅供参考，最终以河北省教育考试院官方公布为准。',
    ]
    
    # 控制线从plans.json获取
    with open(os.path.join(RAW_DIR, 'plans.json'), 'r', encoding='utf-8') as f:
        plans = json.load(f)
    control_lines = []
    for cid, cat in data.items():
        cl = plans.get(cid, {}).get('control_line', '未公布')
        control_lines.append((cat['name'], cl, cat['desc']))
    
    create_info_sheet(wb, '2025年河北单招各类一分一档表', 'http://www.hbdzxx.com/score/yifenyidang/2025/', control_lines, notes)
    
    for cid, cat in data.items():
        sheet_name = get_sheet_name(cat['name'])
        ws = wb.create_sheet(sheet_name)
        
        # 顶部标注
        total_candidates = cat['data'][-1][2] if cat['data'] else '未知'
        ws.cell(row=1, column=1, value=f'{cat["name"]}（{cat["desc"]}）').font = TITLE_FONT
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        ws.cell(row=2, column=1, value=f'总考生人数：{total_candidates}人').font = Font(name='微软雅黑', size=11, bold=True, color='FF0000')
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=5)
        ws.cell(row=3, column=1, value=f'数据来源：{cat["source_url"]}').font = NOTE_FONT
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=5)
        
        # 表头
        headers = ['序号', '分数', '本段人数', '累计人数', '累计占比']
        for col, h in enumerate(headers, 1):
            ws.cell(row=5, column=col, value=h)
        style_header_row(ws, 5, len(headers))
        
        # 数据
        for i, row_data in enumerate(cat['data'], 1):
            ws.cell(row=5 + i, column=1, value=i)
            ws.cell(row=5 + i, column=2, value=row_data[0])
            try:
                ws.cell(row=5 + i, column=3, value=int(row_data[1]))
            except:
                ws.cell(row=5 + i, column=3, value=row_data[1])
            try:
                ws.cell(row=5 + i, column=4, value=int(row_data[2]))
            except:
                ws.cell(row=5 + i, column=4, value=row_data[2])
            ws.cell(row=5 + i, column=5, value=row_data[3])
        
        style_data_cells(ws, 6, 5 + len(cat['data']), len(headers))
        auto_column_width(ws, len(headers))
        ws.freeze_panes = 'A6'
        
        print(f'  {sheet_name}: {len(cat["data"])} 条')
    
    filepath = os.path.join(OUT_DIR, '2025河北单招_一分一档.xlsx')
    wb.save(filepath)
    print(f'  保存: {filepath}')
    return filepath

def generate_school_details():
    """生成院校详情工作簿"""
    print('=== 生成院校详情工作簿 ===')
    details_file = os.path.join(RAW_DIR, 'school_details.json')
    if not os.path.exists(details_file):
        print('  错误: 找不到院校详情数据')
        return None
    
    with open(details_file, 'r', encoding='utf-8') as f:
        details = json.load(f)
    
    wb = Workbook()
    wb.remove(wb.active)
    
    # 统计
    public_count = sum(1 for d in details.values() if d.get('办学性质') == '公办')
    private_count = sum(1 for d in details.values() if d.get('办学性质') == '民办')
    
    notes = [
        f'1. 本表汇总了2025年河北单招十大类招生涉及的全部 {len(details)} 所院校（已去重）。',
        f'2. 其中公办院校 {public_count} 所，民办院校 {private_count} 所。',
        '3. "详细地址"列标注"详见院校官网"的，表示河北单招网未提供该院校的详细通讯地址，请查阅院校官方网站。',
        '4. "院校层次"为本科/专科，本科院校同时招收高职单招专科层次学生。',
        '5. "招生大类数"为该院校在2025年单招中涉及的考试类数量（以院校详情页展示为准）。',
        '6. 数据来源于河北单招网各院校详情页，仅供参考。',
    ]
    
    create_info_sheet(wb, '2025年河北单招院校详情（含地址）', 'http://www.hbdzxx.com/school/', None, notes)
    
    # 所有院校一个sheet，按办学性质排序
    ws = wb.create_sheet('院校详情')
    
    headers = ['序号', '院校名称', '办学性质', '所在城市', '院校层次', '院校代码', '招生大类数', '详细地址', '详情页URL']
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header_row(ws, 1, len(headers))
    
    # 排序：公办在前，然后按名称
    sorted_schools = sorted(details.items(), key=lambda x: (0 if x[1].get('办学性质') == '公办' else 1, x[0]))
    
    for i, (name, detail) in enumerate(sorted_schools, 1):
        ws.cell(row=1 + i, column=1, value=i)
        ws.cell(row=1 + i, column=2, value=name)
        ws.cell(row=1 + i, column=3, value=detail.get('办学性质', ''))
        ws.cell(row=1 + i, column=4, value=detail.get('所在城市', ''))
        ws.cell(row=1 + i, column=5, value=detail.get('院校层次', ''))
        ws.cell(row=1 + i, column=6, value=detail.get('院校代码', ''))
        ws.cell(row=1 + i, column=7, value=detail.get('招生大类数', ''))
        ws.cell(row=1 + i, column=8, value=detail.get('详细地址', ''))
        ws.cell(row=1 + i, column=9, value=detail.get('详情页URL', ''))
    
    style_data_cells(ws, 2, 1 + len(details), len(headers), left_cols=[2, 8, 9])
    auto_column_width(ws, len(headers), max_width=60)
    ws.freeze_panes = 'A2'
    
    print(f'  院校详情: {len(details)} 所 (公办{public_count}, 民办{private_count})')
    
    filepath = os.path.join(OUT_DIR, '2025河北单招_院校详情_含地址.xlsx')
    wb.save(filepath)
    print(f'  保存: {filepath}')
    return filepath

if __name__ == '__main__':
    f1 = generate_score_lines()
    f2 = generate_plans()
    f3 = generate_yifenyidang()
    f4 = generate_school_details()
    print('\n=== 全部生成完成 ===')
    for f in [f1, f2, f3, f4]:
        if f:
            print(f'  {f}')
