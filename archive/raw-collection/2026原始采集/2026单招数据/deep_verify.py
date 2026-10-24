import json
import os
import openpyxl

base_dir = r"C:\Users\XIAO\Desktop\大类信息\2026单招数据"
categories = ['一类','二类','三类','四类','五类','六类','七类','八类美术','九类','十类']

print("=" * 70)
print("深度细节核对：逐条数据验证")
print("=" * 70)

# 1. 一分一档累计人数逻辑验证
print("\n【1】一分一档累计人数逻辑验证")
for cat in categories:
    yfyd_path = os.path.join(base_dir, "一分一档", f"2026年河北单招_{cat}_一分一档.json")
    with open(yfyd_path, 'r', encoding='utf-8') as f:
        yfyd = json.load(f)
    
    errors = 0
    cumulative = 0
    for i, row in enumerate(yfyd['data']):
        try:
            this_count = int(row.get('本段人数', 0))
            cum_count = int(row.get('累计人数', 0))
            cumulative += this_count
            if cumulative != cum_count:
                errors += 1
                if errors <= 3:
                    print(f"  ❌ {cat} 第{i+1}行: 计算累计={cumulative} 表中累计={cum_count}")
        except:
            pass
    
    if errors == 0:
        print(f"  ✓ {cat}: {len(yfyd['data'])}段，累计逻辑全部正确，总人数{cumulative}")
    else:
        print(f"  ❌ {cat}: 发现{errors}处累计错误")

# 2. 分数线与招生计划专业对应验证
print("\n【2】分数线与招生计划专业对应验证")
for cat in categories:
    score_path = os.path.join(base_dir, "分数线", f"2026年河北单招_{cat}_分数线.json")
    plan_path = os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json")
    
    with open(score_path, 'r', encoding='utf-8') as f:
        score_data = json.load(f)
    with open(plan_path, 'r', encoding='utf-8') as f:
        plan_data = json.load(f)
    
    score_keys = set()
    for row in score_data['data']:
        score_keys.add((row.get('院校名称',''), row.get('专业名称','')))
    
    plan_keys = set()
    for row in plan_data['data']:
        plan_keys.add((row.get('院校名称',''), row.get('专业名称','')))
    
    only_in_score = score_keys - plan_keys
    only_in_plan = plan_keys - score_keys
    
    if len(score_keys) == len(plan_keys) and not only_in_score and not only_in_plan:
        print(f"  ✓ {cat}: 分数线{len(score_keys)}条 = 招生计划{len(plan_keys)}条，完全对应")
    else:
        print(f"  ⚠️ {cat}: 分数线{len(score_keys)}条, 招生计划{len(plan_keys)}条")
        if only_in_score:
            print(f"     仅在分数线中: {len(only_in_score)}条")
            for k in list(only_in_score)[:3]:
                print(f"       - {k[0]} - {k[1]}")
        if only_in_plan:
            print(f"     仅在招生计划中: {len(only_in_plan)}条")
            for k in list(only_in_plan)[:3]:
                print(f"       - {k[0]} - {k[1]}")

# 3. 院校详情中院校名称与招生计划对应验证
print("\n【3】院校详情与招生计划院校对应验证")
for cat in categories:
    school_path = os.path.join(base_dir, "院校详情", f"2026年河北单招_{cat}_院校详情.json")
    plan_path = os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json")
    
    with open(school_path, 'r', encoding='utf-8') as f:
        school_data = json.load(f)
    with open(plan_path, 'r', encoding='utf-8') as f:
        plan_data = json.load(f)
    
    school_names = set(row.get('院校名称','') for row in school_data['data'])
    plan_schools = set(row.get('院校名称','') for row in plan_data['data'])
    
    if school_names == plan_schools:
        print(f"  ✓ {cat}: 院校详情{len(school_names)}所 = 招生计划{len(plan_schools)}所，完全对应")
    else:
        missing_in_school = plan_schools - school_names
        missing_in_plan = school_names - plan_schools
        print(f"  ⚠️ {cat}: 院校详情{len(school_names)}所, 招生计划{len(plan_schools)}所")
        if missing_in_school:
            print(f"     招生计划有但院校详情没有: {len(missing_in_school)}所")
            for s in list(missing_in_school)[:3]:
                print(f"       - {s}")
        if missing_in_plan:
            print(f"     院校详情有但招生计划没有: {len(missing_in_plan)}所")
            for s in list(missing_in_plan)[:3]:
                print(f"       - {s}")

# 4. 增强版工作簿具体数据抽查
print("\n【4】增强版工作簿具体数据抽查（每类抽3条）")
import random
random.seed(42)

for cat in categories:
    major_path = os.path.join(base_dir, "专业计划及学费", f"2026年河北单招_{cat}_专业计划及学费.xlsx")
    plan_path = os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json")
    
    with open(plan_path, 'r', encoding='utf-8') as f:
        plan_data = json.load(f)
    
    wb = openpyxl.load_workbook(major_path)
    ws = wb.active
    excel_rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[3]:
            excel_rows.append(row)
    
    # 随机抽3条
    samples = random.sample(range(len(excel_rows)), min(3, len(excel_rows)))
    all_match = True
    for idx in samples:
        row = excel_rows[idx]
        school = row[1]
        major = row[3]
        plan_num = row[4]
        xuefei = row[6]
        
        # 在JSON中查找
        found = False
        for p in plan_data['data']:
            if p.get('院校名称','') == school and p.get('专业名称','') == major:
                found = True
                if str(plan_num) != str(p.get('计划数','')) or str(xuefei) != str(p.get('学费(元/年)','')):
                    all_match = False
                    print(f"  ❌ {cat} 抽查: {school}-{major}")
                    print(f"     Excel: 计划={plan_num}, 学费={xuefei}")
                    print(f"     JSON:  计划={p.get('计划数','')}, 学费={p.get('学费(元/年)','')}")
                break
        
        if not found:
            all_match = False
            print(f"  ❌ {cat} 抽查: {school}-{major} 在JSON中找不到")
    
    if all_match:
        print(f"  ✓ {cat}: 抽查3条全部匹配")

# 5. 院校地址格式检查
print("\n【5】院校地址格式检查")
addr_issues = []
for cat in categories:
    school_path = os.path.join(base_dir, "院校详情_增强版", f"2026年河北单招_{cat}_院校详情_含地址.xlsx")
    wb = openpyxl.load_workbook(school_path)
    ws = wb.active
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[1]:
            addr = row[4] or ''
            if not addr or addr == '待补充':
                addr_issues.append(f"{cat}: {row[1]} 地址为空")
            elif len(addr) < 5:
                addr_issues.append(f"{cat}: {row[1]} 地址过短: {addr}")
            elif not any(kw in addr for kw in ['省','市','区','县','路','街','号','大道','镇','乡']):
                addr_issues.append(f"{cat}: {row[1]} 地址格式可疑: {addr}")

if addr_issues:
    print(f"  ⚠️ 发现{len(addr_issues)}个地址问题:")
    for issue in addr_issues[:10]:
        print(f"     - {issue}")
else:
    print("  ✓ 所有院校地址格式正常")

# 6. 学费数据合理性检查
print("\n【6】学费数据合理性检查")
xuefei_issues = []
for cat in categories:
    plan_path = os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json")
    with open(plan_path, 'r', encoding='utf-8') as f:
        plan_data = json.load(f)
    
    for row in plan_data['data']:
        xuefei = row.get('学费(元/年)', '')
        try:
            xf = int(xuefei)
            if xf < 3000 or xf > 30000:
                xuefei_issues.append(f"{cat}: {row.get('院校名称','')}-{row.get('专业名称','')} 学费={xuefei}")
        except:
            if xuefei and '暂定' not in xuefei and '以' not in xuefei:
                xuefei_issues.append(f"{cat}: {row.get('院校名称','')}-{row.get('专业名称','')} 学费非数字: {xuefei}")

if xuefei_issues:
    print(f"  ⚠️ 发现{len(xuefei_issues)}个学费异常:")
    for issue in xuefei_issues[:10]:
        print(f"     - {issue}")
else:
    print("  ✓ 所有学费数据在合理范围内")

print("\n" + "=" * 70)
print("深度核对完成")
print("=" * 70)
