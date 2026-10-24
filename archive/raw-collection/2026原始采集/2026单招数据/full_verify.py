import json
import os
import openpyxl

base_dir = r"C:\Users\XIAO\Desktop\大类信息\2026单招数据"
categories = ['一类','二类','三类','四类','五类','六类','七类','八类美术','九类','十类']

print("=" * 70)
print("第一阶段：原始数据与增强版数据一致性核对")
print("=" * 70)

errors = []
warnings = []

for cat in categories:
    print(f"\n--- {cat} ---")
    
    # 1. 读取原始招生计划JSON
    plan_json_path = os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json")
    with open(plan_json_path, 'r', encoding='utf-8') as f:
        plan_json = json.load(f)
    plan_json_count = len(plan_json['data'])
    
    # 2. 读取原始分数线JSON
    score_json_path = os.path.join(base_dir, "分数线", f"2026年河北单招_{cat}_分数线.json")
    with open(score_json_path, 'r', encoding='utf-8') as f:
        score_json = json.load(f)
    score_json_count = len(score_json['data'])
    
    # 3. 读取原始院校详情JSON
    school_json_path = os.path.join(base_dir, "院校详情", f"2026年河北单招_{cat}_院校详情.json")
    with open(school_json_path, 'r', encoding='utf-8') as f:
        school_json = json.load(f)
    school_json_count = len(school_json['data'])
    
    # 4. 读取增强版院校详情Excel
    school_xlsx_path = os.path.join(base_dir, "院校详情_增强版", f"2026年河北单招_{cat}_院校详情_含地址.xlsx")
    wb = openpyxl.load_workbook(school_xlsx_path)
    ws = wb.active
    school_xlsx_rows = []
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[1]:
            school_xlsx_rows.append(row)
    school_xlsx_count = len(school_xlsx_rows)
    
    # 5. 读取专业计划及学费Excel
    major_xlsx_path = os.path.join(base_dir, "专业计划及学费", f"2026年河北单招_{cat}_专业计划及学费.xlsx")
    wb2 = openpyxl.load_workbook(major_xlsx_path)
    ws2 = wb2.active
    major_xlsx_rows = []
    for row in ws2.iter_rows(min_row=4, values_only=True):
        if row[3]:
            major_xlsx_rows.append(row)
    major_xlsx_count = len(major_xlsx_rows)
    
    # 核对行数
    print(f"  原始招生计划JSON: {plan_json_count} 条")
    print(f"  专业计划Excel: {major_xlsx_count} 条")
    if plan_json_count != major_xlsx_count:
        errors.append(f"{cat}: 招生计划行数不一致 JSON={plan_json_count} Excel={major_xlsx_count}")
        print(f"  ❌ 行数不一致!")
    else:
        print(f"  ✓ 行数一致")
    
    print(f"  原始分数线JSON: {score_json_count} 条")
    print(f"  原始院校详情JSON: {school_json_count} 条")
    print(f"  增强版院校详情Excel: {school_xlsx_count} 条")
    if school_json_count != school_xlsx_count:
        errors.append(f"{cat}: 院校详情行数不一致 JSON={school_json_count} Excel={school_xlsx_count}")
        print(f"  ❌ 院校详情行数不一致!")
    else:
        print(f"  ✓ 院校详情行数一致")
    
    # 6. 核对专业计划数据内容
    print(f"  核对专业计划数据内容...")
    plan_json_dict = {}
    for row in plan_json['data']:
        key = (row.get('院校名称',''), row.get('专业名称',''))
        plan_json_dict[key] = row
    
    content_errors = 0
    for row in major_xlsx_rows:
        school = row[1]
        major = row[3]
        plan_num = row[4]
        xuezhi = row[5]
        xuefei = row[6]
        score = row[7]
        
        key = (school, major)
        if key in plan_json_dict:
            orig = plan_json_dict[key]
            if str(plan_num) != str(orig.get('计划数','')):
                content_errors += 1
                if content_errors <= 3:
                    errors.append(f"{cat}: 计划数不一致 {school}-{major} Excel={plan_num} JSON={orig.get('计划数','')}")
            if str(xuezhi) != str(orig.get('学制','')):
                content_errors += 1
                if content_errors <= 3:
                    errors.append(f"{cat}: 学制不一致 {school}-{major} Excel={xuezhi} JSON={orig.get('学制','')}")
            if str(xuefei) != str(orig.get('学费(元/年)','')):
                content_errors += 1
                if content_errors <= 3:
                    errors.append(f"{cat}: 学费不一致 {school}-{major} Excel={xuefei} JSON={orig.get('学费(元/年)','')}")
        else:
            content_errors += 1
            if content_errors <= 3:
                errors.append(f"{cat}: Excel中专业在JSON中找不到 {school}-{major}")
    
    if content_errors == 0:
        print(f"  ✓ 专业计划数据内容全部一致")
    else:
        print(f"  ❌ 发现 {content_errors} 处内容不一致")
    
    # 7. 核对分数线关联
    print(f"  核对分数线关联...")
    score_json_dict = {}
    for row in score_json['data']:
        key = (row.get('院校名称',''), row.get('专业名称',''))
        score_json_dict[key] = row.get('录取投档分','')
    
    score_mismatch = 0
    score_missing = 0
    for row in major_xlsx_rows:
        school = row[1]
        major = row[3]
        score = row[7]
        key = (school, major)
        if key in score_json_dict:
            if str(score) != str(score_json_dict[key]):
                score_mismatch += 1
                if score_mismatch <= 3:
                    errors.append(f"{cat}: 分数线不一致 {school}-{major} Excel={score} JSON={score_json_dict[key]}")
        else:
            score_missing += 1
    
    if score_mismatch == 0 and score_missing == 0:
        print(f"  ✓ 分数线全部正确关联")
    else:
        print(f"  ⚠️ 分数线不匹配={score_mismatch}, 缺失={score_missing}")
        if score_missing > 0:
            warnings.append(f"{cat}: {score_missing}个专业在分数线中找不到对应记录（可能是该专业无投档数据）")
    
    # 8. 核对院校地址
    print(f"  核对院校地址...")
    addr_missing = 0
    for row in school_xlsx_rows:
        addr = row[4] or ''
        if not addr or addr == '待补充':
            addr_missing += 1
            errors.append(f"{cat}: 地址缺失 {row[1]}")
    
    if addr_missing == 0:
        print(f"  ✓ 院校地址全部填充")
    else:
        print(f"  ❌ {addr_missing} 所院校地址缺失")
    
    # 9. 核对院校详情中的专业数和计划总数
    print(f"  核对院校详情统计数据...")
    school_stats = {}
    for row in plan_json['data']:
        school = row.get('院校名称','')
        if school not in school_stats:
            school_stats[school] = {'count': 0, 'total': 0}
        school_stats[school]['count'] += 1
        try:
            school_stats[school]['total'] += int(row.get('计划数', 0))
        except:
            pass
    
    stat_errors = 0
    for row in school_xlsx_rows:
        school = row[1]
        major_count = row[5]
        total_plan = row[6]
        if school in school_stats:
            if major_count != school_stats[school]['count']:
                stat_errors += 1
                if stat_errors <= 3:
                    errors.append(f"{cat}: 专业数不一致 {school} Excel={major_count} 计算={school_stats[school]['count']}")
            if total_plan != school_stats[school]['total']:
                stat_errors += 1
                if stat_errors <= 3:
                    errors.append(f"{cat}: 计划总数不一致 {school} Excel={total_plan} 计算={school_stats[school]['total']}")
    
    if stat_errors == 0:
        print(f"  ✓ 院校详情统计数据正确")
    else:
        print(f"  ❌ 发现 {stat_errors} 处统计错误")

print("\n" + "=" * 70)
print("第二阶段：数据汇总核对")
print("=" * 70)

# 总计划数核对
total_plan_json = 0
total_plan_excel = 0
for cat in categories:
    plan_json_path = os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json")
    with open(plan_json_path, 'r', encoding='utf-8') as f:
        plan_json = json.load(f)
    for row in plan_json['data']:
        try:
            total_plan_json += int(row.get('计划数', 0))
        except:
            pass
    
    major_xlsx_path = os.path.join(base_dir, "专业计划及学费", f"2026年河北单招_{cat}_专业计划及学费.xlsx")
    wb = openpyxl.load_workbook(major_xlsx_path)
    ws = wb.active
    for row in ws.iter_rows(min_row=4, values_only=True):
        if row[3]:
            try:
                total_plan_excel += int(row[4] or 0)
            except:
                pass

print(f"JSON总计划数: {total_plan_json}")
print(f"Excel总计划数: {total_plan_excel}")
if total_plan_json == total_plan_excel:
    print("✓ 总计划数一致")
else:
    errors.append(f"总计划数不一致 JSON={total_plan_json} Excel={total_plan_excel}")
    print("❌ 总计划数不一致")

# 一分一档总人数核对
print("\n一分一档总人数核对:")
for cat in categories:
    yfyd_path = os.path.join(base_dir, "一分一档", f"2026年河北单招_{cat}_一分一档.json")
    with open(yfyd_path, 'r', encoding='utf-8') as f:
        yfyd = json.load(f)
    if yfyd['data']:
        last_row = yfyd['data'][-1]
        total = last_row.get('累计人数', 0)
        print(f"  {cat}: {total} 人")

print("\n" + "=" * 70)
print("验证结果汇总")
print("=" * 70)
print(f"错误数量: {len(errors)}")
print(f"警告数量: {len(warnings)}")

if errors:
    print("\n❌ 发现的错误:")
    for e in errors[:30]:
        print(f"  - {e}")
    if len(errors) > 30:
        print(f"  ... 还有 {len(errors)-30} 条错误")
else:
    print("\n✓ 未发现数据错误")

if warnings:
    print("\n⚠️ 警告信息:")
    for w in warnings[:10]:
        print(f"  - {w}")
