import json
import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# 基础目录
base_dir = r"C:\Users\XIAO\Desktop\大类信息\2026单招数据"

# 数据类型配置
data_types = {
    "分数线": {
        "dir": "分数线",
        "sheet_name": "录取分数线",
        "title": "2026年河北高职单招录取分数线"
    },
    "招生计划": {
        "dir": "招生计划",
        "sheet_name": "招生计划",
        "title": "2026年河北高职单招招生计划"
    },
    "院校详情": {
        "dir": "院校详情",
        "sheet_name": "招生院校",
        "title": "2026年河北高职单招招生院校详情"
    },
    "一分一档": {
        "dir": "一分一档",
        "sheet_name": "一分一档表",
        "title": "2026年河北高职单招一分一档表"
    }
}

# 类别名称映射（用于标题）
category_full_names = {
    "一类": "考试一类（土木建筑、资源环境与安全、水利等）",
    "二类": "考试二类（交通运输、能源动力与材料等）",
    "三类": "考试三类（装备制造等）",
    "四类": "考试四类（农林牧渔、轻工纺织、生物与化工等）",
    "五类": "考试五类（公共管理与服务、旅游、公安与司法等）",
    "六类": "考试六类（教育与体育等）",
    "七类": "考试七类（医药卫生等）",
    "八类美术": "考试八类-美术设计类",
    "九类": "考试九类（财经商贸等）",
    "十类": "考试十类（电子与信息等）"
}

# 样式定义
header_font = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

title_font = Font(name='微软雅黑', size=14, bold=True)
title_alignment = Alignment(horizontal='center', vertical='center')

data_font = Font(name='微软雅黑', size=10)
data_alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

thin_border = Border(
    left=Side(style='thin', color='B4B4B4'),
    right=Side(style='thin', color='B4B4B4'),
    top=Side(style='thin', color='B4B4B4'),
    bottom=Side(style='thin', color='B4B4B4')
)

# 交替行颜色
even_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')

def convert_json_to_excel(json_path, excel_path, data_type_config, category_name):
    """将单个JSON文件转换为Excel"""
    with open(json_path, 'r', encoding='utf-8') as f:
        content = json.load(f)
    
    headers = content['headers']
    data = content['data']
    
    wb = Workbook()
    ws = wb.active
    ws.title = data_type_config['sheet_name']
    
    # 写入标题行
    full_title = f"{data_type_config['title']} - {category_full_names.get(category_name, category_name)}"
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=full_title)
    title_cell.font = title_font
    title_cell.alignment = title_alignment
    ws.row_dimensions[1].height = 30
    
    # 写入数据来源说明
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers))
    source_cell = ws.cell(row=2, column=1, value=f"数据来源：河北省教育考试院 | 共 {len(data)} 条记录 | 提取自 hbdzxx.com")
    source_cell.font = Font(name='微软雅黑', size=9, italic=True, color='666666')
    source_cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 20
    
    # 写入表头（第3行）
    header_row = 3
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border
    ws.row_dimensions[header_row].height = 25
    
    # 写入数据
    for row_idx, row_data in enumerate(data, header_row + 1):
        for col_idx, header in enumerate(headers, 1):
            value = row_data.get(header, '')
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.border = thin_border
            
            # 数字列居中，文本列左对齐
            if header in ['录取投档分', '计划数', '学制', '学费(元/年)', '分数', '本段人数', '累计人数', '累计占比', '序号']:
                cell.alignment = center_alignment
            else:
                cell.alignment = data_alignment
            
            # 交替行颜色
            if (row_idx - header_row) % 2 == 0:
                cell.fill = even_fill
    
    # 设置列宽
    col_widths = {
        '院校名称': 28,
        '专业名称': 45,
        '考试类': 12,
        '录取投档分': 12,
        '计划数': 10,
        '学制': 8,
        '学费(元/年)': 14,
        '序号': 8,
        '所在地区': 12,
        '办学性质': 12,
        '操作': 15,
        '分数': 10,
        '本段人数': 12,
        '累计人数': 12,
        '累计占比': 12,
    }
    
    for col_idx, header in enumerate(headers, 1):
        width = col_widths.get(header, 20)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    
    # 冻结窗格（冻结标题和表头）
    ws.freeze_panes = f'A{header_row + 1}'
    
    # 添加筛选
    ws.auto_filter.ref = f'A{header_row}:{get_column_letter(len(headers))}{header_row + len(data)}'
    
    wb.save(excel_path)
    return len(data)

# 主转换流程
print("=" * 60)
print("开始将JSON数据转换为Excel工作簿")
print("=" * 60)

total_files = 0
total_records = 0

for type_name, type_config in data_types.items():
    input_dir = os.path.join(base_dir, type_config['dir'])
    output_dir = os.path.join(base_dir, type_config['dir'])
    
    print(f"\n【{type_name}】")
    
    for filename in sorted(os.listdir(input_dir)):
        if not filename.endswith('.json'):
            continue
        
        json_path = os.path.join(input_dir, filename)
        
        # 从文件名提取类别名称
        # 格式: 2026年河北单招_{类别}_分数线.json
        parts = filename.replace('.json', '').split('_')
        if len(parts) >= 2:
            category_name = parts[1]
        else:
            category_name = "未知"
        
        excel_filename = filename.replace('.json', '.xlsx')
        excel_path = os.path.join(output_dir, excel_filename)
        
        count = convert_json_to_excel(json_path, excel_path, type_config, category_name)
        total_files += 1
        total_records += count
        
        print(f"  ✓ {excel_filename} ({count} 条记录)")

print("\n" + "=" * 60)
print(f"转换完成！共生成 {total_files} 个Excel工作簿，总计 {total_records} 条记录")
print("=" * 60)
