import json
import os

base_dir = r"C:\Users\XIAO\Desktop\大类信息\2026单招数据"
output_dir = r"C:\Users\XIAO\Desktop\大类信息\Admission_Info"

categories = ['一类','二类','三类','四类','五类','六类','七类','八类美术','九类','十类']
cat_ids = {'一类':'1','二类':'2','三类':'3','四类':'4','五类':'5','六类':'6','七类':'7','八类美术':'8','九类':'9','十类':'10'}

# 读取所有数据
all_data = {}

for cat in categories:
    cat_id = cat_ids[cat]
    all_data[cat_id] = {'name': cat}
    
    # 分数线
    with open(os.path.join(base_dir, "分数线", f"2026年河北单招_{cat}_分数线.json"), 'r', encoding='utf-8') as f:
        all_data[cat_id]['scoreLines'] = json.load(f)['data']
    
    # 招生计划
    with open(os.path.join(base_dir, "招生计划", f"2026年河北单招_{cat}_招生计划.json"), 'r', encoding='utf-8') as f:
        all_data[cat_id]['plans'] = json.load(f)['data']
    
    # 院校详情
    with open(os.path.join(base_dir, "院校详情", f"2026年河北单招_{cat}_院校详情.json"), 'r', encoding='utf-8') as f:
        all_data[cat_id]['schools'] = json.load(f)['data']
    
    # 一分一档
    with open(os.path.join(base_dir, "一分一档", f"2026年河北单招_{cat}_一分一档.json"), 'r', encoding='utf-8') as f:
        all_data[cat_id]['yifenyidang'] = json.load(f)['data']

# 读取院校地址
school_addresses = {}
for cat in categories:
    xp = os.path.join(base_dir, "院校详情_增强版", f"2026年河北单招_{cat}_院校详情_含地址.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(xp)
    ws = wb.active
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[1]:
            school_addresses[r[1]] = r[4] or ''

# 读取专业计划及学费（含地址和分数线）
major_plans = {}
for cat in categories:
    cat_id = cat_ids[cat]
    xp = os.path.join(base_dir, "专业计划及学费", f"2026年河北单招_{cat}_专业计划及学费.xlsx")
    wb = openpyxl.load_workbook(xp)
    ws = wb.active
    major_plans[cat_id] = []
    for r in ws.iter_rows(min_row=4, values_only=True):
        if r[3]:
            major_plans[cat_id].append({
                'school': r[1],
                'address': r[2] or '',
                'major': r[3],
                'plan': r[4],
                'duration': r[5],
                'tuition': r[6],
                'score': r[7]
            })

# 输出JS数据文件
js_content = "const ADMISSION_DATA = {\n"
js_content += "  '2026': {\n"

for cat_id in ['1','2','3','4','5','6','7','8','9','10']:
    cat_name = all_data[cat_id]['name']
    js_content += f"    '{cat_id}': {{\n"
    js_content += f"      name: '{cat_name}',\n"
    js_content += f"      scoreLines: {json.dumps(all_data[cat_id]['scoreLines'], ensure_ascii=False)},\n"
    js_content += f"      plans: {json.dumps(all_data[cat_id]['plans'], ensure_ascii=False)},\n"
    js_content += f"      schools: {json.dumps(all_data[cat_id]['schools'], ensure_ascii=False)},\n"
    js_content += f"      yifenyidang: {json.dumps(all_data[cat_id]['yifenyidang'], ensure_ascii=False)},\n"
    js_content += f"      majorPlans: {json.dumps(major_plans[cat_id], ensure_ascii=False)}\n"
    js_content += f"    }},\n"

js_content += "  },\n"
js_content += "  '2025': null\n"  # 预留2025年数据
js_content += "};\n"

with open(os.path.join(output_dir, "data.js"), 'w', encoding='utf-8') as f:
    f.write(js_content)

print(f"数据导出完成")
print(f"分数线: {sum(len(all_data[c]['scoreLines']) for c in all_data)}条")
print(f"招生计划: {sum(len(all_data[c]['plans']) for c in all_data)}条")
print(f"院校详情: {sum(len(all_data[c]['schools']) for c in all_data)}所次")
print(f"一分一档: {sum(len(all_data[c]['yifenyidang']) for c in all_data)}段")
print(f"专业计划: {sum(len(major_plans[c]) for c in major_plans)}条")
print(f"文件大小: {os.path.getsize(os.path.join(output_dir, 'data.js')) / 1024:.1f} KB")
