# -*- coding: utf-8 -*-
"""
将2024和2025年Excel数据转换为JS格式
"""
import openpyxl
import json
import os

BASE = r'C:\Users\XIAO\Desktop\大类信息\Admission_Info\data'
OUT_DIR = r'C:\Users\XIAO\Desktop\大类信息\Admission_Info'

CATEGORIES = [
    ('1', '一类', '考试一类', '土木建筑、资源环境与安全、水利等'),
    ('2', '二类', '考试二类', '交通运输、能源动力与材料等'),
    ('3', '三类', '考试三类', '装备制造等'),
    ('4', '四类', '考试四类', '农林牧渔、轻工纺织、生物与化工、食品药品与粮食等'),
    ('5', '五类', '考试五类', '公共管理与服务、旅游、公安与司法等'),
    ('6', '六类', '考试六类', '教育与体育、新闻传播等'),
    ('7', '七类', '考试七类', '医药卫生（临床医学类专业除外）等'),
    ('8', '八类美术', '考试八类-美术设计类', '文化艺术等'),
    ('9', '九类', '考试九类', '财经商贸等'),
    ('10', '十类', '考试十类', '电子信息、计算机等'),
]

TYPES = ['分数线', '招生计划', '院校详情', '一分一档', '专业计划及学费']
TYPE_KEYS = {
    '分数线': 'scoreLines',
    '招生计划': 'plans',
    '院校详情': 'schools',
    '一分一档': 'yifenyidang',
    '专业计划及学费': 'majors',
}


def read_excel_data(filepath, header_row=3):
    if not os.path.exists(filepath):
        return []
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    headers = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=col).value
        if v:
            headers.append(str(v).strip())
    data = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in row):
            continue
        row_data = {}
        for i, h in enumerate(headers):
            val = row[i] if i < len(row) else None
            row_data[h] = str(val).strip() if isinstance(val, str) else (val if val is not None else '')
        data.append(row_data)
    return data


def convert_year(year):
    result = {}
    year_dir = os.path.join(BASE, year)
    for cat_id, cat_name, cat_full, cat_desc in CATEGORIES:
        print(f'  [{year}] {cat_name} ...')
        cat_data = {'id': cat_id, 'name': cat_name, 'fullName': cat_full, 'desc': cat_desc}
        for t in TYPES:
            key = TYPE_KEYS[t]
            filepath = os.path.join(year_dir, t, f'{year}年河北单招_{cat_name}_{t}.xlsx')
            cat_data[key] = read_excel_data(filepath)
        result[cat_id] = cat_data
        counts = {t: len(cat_data[TYPE_KEYS[t]]) for t in TYPES}
        print(f'    {counts}')

    out_file = os.path.join(OUT_DIR, f'data_{year}.js')
    js_content = f'const DANZHAO_DATA_{year} = ' + json.dumps(result, ensure_ascii=False, separators=(',', ':')) + ';\n'
    with open(out_file, 'w', encoding='utf-8') as f:
        f.write(js_content)
    size_kb = os.path.getsize(out_file) / 1024
    print(f'  输出: {out_file} ({size_kb:.1f} KB)')


if __name__ == '__main__':
    for year in ['2024', '2025']:
        print(f'===== 转换 {year} 年 =====')
        convert_year(year)
    print('\n全部完成！')
