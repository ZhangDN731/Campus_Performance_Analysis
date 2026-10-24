# -*- coding: utf-8 -*-
"""
将2025年单招数据（3个汇总文件）转换为2026年标准格式（分目录单文件）
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os

BASE = r'C:\Users\XIAO\Desktop\大类信息'
OUT_DIR = os.path.join(BASE, '2025单招数据')

CATEGORIES = [
    ('一类', '一类', '考试一类（土木建筑、资源环境与安全、水利等）'),
    ('二类', '二类', '考试二类（交通运输、能源动力与材料等）'),
    ('三类', '三类', '考试三类（装备制造等）'),
    ('四类', '四类', '考试四类（农林牧渔、轻工纺织、生物与化工、食品药品与粮食等）'),
    ('五类', '五类', '考试五类（公共管理与服务、旅游、公安与司法等）'),
    ('六类', '六类', '考试六类（教育与体育、新闻传播等）'),
    ('七类', '七类', '考试七类（医药卫生（临床医学类专业除外）等）'),
    ('八类美术', '八类美术', '考试八类-美术设计类（文化艺术等）'),
    ('九类', '九类', '考试九类（财经商贸等）'),
    ('十类', '十类', '考试十类（电子信息、计算机等）'),
]

HEADER_FILL = PatternFill(start_color='0891B2', end_color='0891B2', fill_type='solid')
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
TITLE_FONT = Font(name='微软雅黑', size=14, bold=True, color='0E7490')
SUBTITLE_FONT = Font(name='微软雅黑', size=10, color='64748B')
DATA_FONT = Font(name='微软雅黑', size=10)
ALT_FILL = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0'),
)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)


def create_workbook(title, subtitle, headers, data_rows, col_widths=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31] if len(title) > 31 else title
    ncols = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = TITLE_FONT
    cell.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 30
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    cell = ws.cell(row=2, column=1, value=subtitle)
    cell.font = SUBTITLE_FONT
    cell.alignment = CENTER_ALIGN
    ws.row_dimensions[2].height = 20
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[3].height = 25
    for row_idx, row_data in enumerate(data_rows, 4):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col_idx == 1 or (isinstance(val, (int, float)) and not isinstance(val, bool)):
                cell.alignment = CENTER_ALIGN
            else:
                cell.alignment = LEFT_ALIGN
            if row_idx % 2 == 0:
                cell.fill = ALT_FILL
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A4'
    ws.auto_filter.ref = f'A3:{get_column_letter(ncols)}{3 + len(data_rows)}'
    return wb


def extract_region_from_name(name):
    """从院校名称中提取所在地区"""
    if not name:
        return '未知'
    name = str(name)
    cities = ['石家庄', '唐山', '秦皇岛', '邯郸', '邢台', '保定', '张家口', '承德', '沧州', '廊坊', '衡水']
    for c in cities:
        if c in name:
            return c
    if '北京' in name:
        return '北京'
    if '天津' in name:
        return '天津'
    if '曹妃甸' in name:
        return '唐山'
    if '渤海' in name:
        return '沧州'
    return '省内'


def infer_nature(fee):
    """从学费推断办学性质"""
    try:
        f = int(fee)
        if f <= 8000:
            return '公办'
        else:
            return '民办'
    except (ValueError, TypeError):
        return '未知'


def read_sheet_data(wb, sheet_name, header_row=5):
    """读取汇总文件中某个sheet的数据，跳过前header_row-1行说明"""
    if sheet_name not in wb.sheetnames:
        return [], None
    ws = wb[sheet_name]
    # 读取省控制线（第2行）
    control_line = None
    row2 = ws.cell(row=2, column=1).value
    if row2 and '控制线' in str(row2):
        m = str(row2).replace('省控制线：', '').replace('分', '').strip()
        try:
            control_line = int(m)
        except ValueError:
            pass
    # 读取表头
    headers = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=col).value
        if v:
            headers.append(v)
    # 读取数据
    data = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if row[0] is None and row[1] is None:
            continue
        data.append(row[:len(headers)])
    return data, control_line


def process_2025():
    subdirs = ['分数线', '招生计划', '一分一档', '院校详情', '专业计划及学费']
    for sd in subdirs:
        os.makedirs(os.path.join(OUT_DIR, sd), exist_ok=True)

    # 加载3个汇总文件
    wb_score = openpyxl.load_workbook(os.path.join(BASE, '2025河北单招_分数线.xlsx'), data_only=True)
    wb_yfyd = openpyxl.load_workbook(os.path.join(BASE, '2025河北单招_一分一档.xlsx'), data_only=True)
    wb_plan = openpyxl.load_workbook(os.path.join(BASE, '2025河北单招_招生计划_含学费.xlsx'), data_only=True)

    total_stats = {'分数线': 0, '招生计划': 0, '一分一档': 0, '院校详情': 0, '专业计划及学费': 0}

    for sheet_name, out_cat, full_name in CATEGORIES:
        print(f'[处理] 2025年 {out_cat} ...')

        # ===== 读取数据 =====
        score_data, control_line = read_sheet_data(wb_score, sheet_name)
        yfyd_data, _ = read_sheet_data(wb_yfyd, sheet_name)
        plan_data, _ = read_sheet_data(wb_plan, sheet_name)

        # ===== 1. 分数线 =====
        score_rows = []
        for row in score_data:
            if len(row) >= 5 and row[1]:
                score_rows.append((row[1], row[2], row[3], row[4]))
        wb_out = create_workbook(
            f'2025年河北高职单招录取分数线 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(score_rows)} 条记录 | 省控制线：{control_line}分',
            ['院校名称', '专业名称', '考试类', '录取投档分'],
            score_rows,
            col_widths=[25, 35, 15, 14]
        )
        wb_out.save(os.path.join(OUT_DIR, '分数线', f'2025年河北单招_{out_cat}_分数线.xlsx'))
        total_stats['分数线'] += 1

        # ===== 2. 招生计划 =====
        plan_rows = []
        for row in plan_data:
            if len(row) >= 7 and row[1]:
                plan_rows.append((row[1], row[2], row[3], row[4], row[5], row[6]))
        total_plan = sum(int(r[3]) for r in plan_rows if r[3] and str(r[3]).isdigit())
        wb_out = create_workbook(
            f'2025年河北高职单招招生计划 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(plan_rows)} 条记录 | 总计划数: {total_plan} 人',
            ['院校名称', '专业名称', '考试类', '计划数', '学制', '学费(元/年)'],
            plan_rows,
            col_widths=[25, 35, 12, 10, 8, 12]
        )
        wb_out.save(os.path.join(OUT_DIR, '招生计划', f'2025年河北单招_{out_cat}_招生计划.xlsx'))
        total_stats['招生计划'] += 1

        # ===== 3. 一分一档 =====
        yfyd_rows = []
        for row in yfyd_data:
            if len(row) >= 5 and row[1]:
                yfyd_rows.append((row[1], row[2], row[3], row[4]))
        total_people = yfyd_rows[-1][2] if yfyd_rows else 0
        wb_out = create_workbook(
            f'2025年河北高职单招一分一档表 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(yfyd_rows)} 条分数段 | 总人数: {total_people} 人',
            ['分数', '本段人数', '累计人数', '累计占比'],
            yfyd_rows,
            col_widths=[15, 12, 12, 12]
        )
        wb_out.save(os.path.join(OUT_DIR, '一分一档', f'2025年河北单招_{out_cat}_一分一档.xlsx'))
        total_stats['一分一档'] += 1

        # ===== 4. 院校详情（从招生计划提取） =====
        school_map = {}  # name -> {region, nature, major_count, plan_total}
        for row in plan_data:
            if len(row) >= 7 and row[1]:
                name = row[1]
                fee = row[6]
                plan_num = row[4]
                if name not in school_map:
                    school_map[name] = {
                        'region': extract_region_from_name(name),
                        'nature': infer_nature(fee),
                        'major_count': 0,
                        'plan_total': 0,
                    }
                school_map[name]['major_count'] += 1
                if plan_num and str(plan_num).isdigit():
                    school_map[name]['plan_total'] += int(plan_num)

        school_rows = []
        for idx, (name, info) in enumerate(school_map.items(), 1):
            school_rows.append((idx, name, info['region'], info['nature'], '招生计划 院校资讯'))
        wb_out = create_workbook(
            f'2025年河北高职单招招生院校详情 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(school_rows)} 所院校 | 办学性质按学费推断',
            ['序号', '院校名称', '所在地区', '办学性质', '操作'],
            school_rows,
            col_widths=[8, 30, 12, 12, 20]
        )
        wb_out.save(os.path.join(OUT_DIR, '院校详情', f'2025年河北单招_{out_cat}_院校详情.xlsx'))
        total_stats['院校详情'] += 1

        # ===== 5. 专业计划及学费（合并招生计划+分数线） =====
        # 建立分数线映射：(院校, 专业) -> 投档分
        score_map = {}
        for row in score_data:
            if len(row) >= 5 and row[1] and row[2]:
                key = (str(row[1]).strip(), str(row[2]).strip())
                score_map[key] = row[4]

        major_rows = []
        for idx, row in enumerate(plan_data, 1):
            if len(row) >= 7 and row[1]:
                key = (str(row[1]).strip(), str(row[2]).strip())
                score = score_map.get(key, '—')
                major_rows.append((
                    idx,
                    row[1],
                    '',  # 2025年无地址信息
                    row[2],
                    row[4],
                    row[5],
                    row[6],
                    score,
                ))
        wb_out = create_workbook(
            f'2025年河北高职单招专业计划及学费 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(major_rows)} 个专业 | 总计划数: {total_plan} 人',
            ['序号', '院校名称', '院校地址', '专业名称', '计划数', '学制', '学费(元/年)', '2025投档最低分'],
            major_rows,
            col_widths=[8, 22, 25, 30, 10, 8, 12, 14]
        )
        wb_out.save(os.path.join(OUT_DIR, '专业计划及学费', f'2025年河北单招_{out_cat}_专业计划及学费.xlsx'))
        total_stats['专业计划及学费'] += 1

        print(f'  ✓ 分数线{len(score_rows)} 招生计划{len(plan_rows)} 一分一档{len(yfyd_rows)} 院校{len(school_rows)} 专业{len(major_rows)}')

    print(f'\n===== 2025年数据转换完成 =====')
    for k, v in total_stats.items():
        print(f'  {k}: {v}个文件')


if __name__ == '__main__':
    process_2025()
