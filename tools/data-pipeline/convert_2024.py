# -*- coding: utf-8 -*-
"""
将2024年单招数据（单文件多sheet）转换为2026年标准格式（分目录单文件）
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import os
import re

BASE = r'C:\Users\XIAO\Desktop\大类信息'
SRC_DIR = BASE
OUT_DIR = os.path.join(BASE, '2024单招数据')

# 类别配置：(源文件名后缀, 输出类别名, 类别全称, 网站类别号)
CATEGORIES = [
    ('一类', '一类', '考试一类（土木建筑、资源环境与安全、水利等）', '1'),
    ('二类', '二类', '考试二类（交通运输、能源动力与材料等）', '2'),
    ('三类', '三类', '考试三类（装备制造等）', '3'),
    ('四类', '四类', '考试四类（农林牧渔、轻工纺织、生物与化工、食品药品与粮食等）', '4'),
    ('五类', '五类', '考试五类（公共管理与服务、旅游、公安与司法等）', '5'),
    ('六类', '六类', '考试六类（教育与体育、新闻传播等）', '6'),
    ('七类', '七类', '考试七类（医药卫生（临床医学类专业除外）等）', '7'),
    ('八类美术设计类', '八类美术', '考试八类-美术设计类（文化艺术等）', '24'),
    ('九类', '九类', '考试九类（财经商贸等）', '9'),
    ('十类', '十类', '考试十类（电子信息、计算机等）', '10'),
]

# 样式定义
HEADER_FILL = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
HEADER_FONT = Font(name='微软雅黑', size=11, bold=True, color='FFFFFF')
TITLE_FONT = Font(name='微软雅黑', size=14, bold=True, color='1E40AF')
SUBTITLE_FONT = Font(name='微软雅黑', size=10, color='64748B')
DATA_FONT = Font(name='微软雅黑', size=10)
ALT_FILL = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
THIN_BORDER = Border(
    left=Side(style='thin', color='E2E8F0'),
    right=Side(style='thin', color='E2E8F0'),
    top=Side(style='thin', color='E2E8F0'),
    bottom=Side(style='thin', color='E2E8F0'),
)
CENTER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
LEFT_ALIGN = Alignment(horizontal='left', vertical='center', wrap_text=True)


def create_workbook(title, subtitle, headers, data_rows, col_widths=None):
    """创建标准格式的工作簿"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31] if len(title) > 31 else title

    ncols = len(headers)
    # 标题行
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = TITLE_FONT
    cell.alignment = CENTER_ALIGN
    ws.row_dimensions[1].height = 30

    # 副标题行
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    cell = ws.cell(row=2, column=1, value=subtitle)
    cell.font = SUBTITLE_FONT
    cell.alignment = CENTER_ALIGN
    ws.row_dimensions[2].height = 20

    # 表头行
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER_ALIGN
        cell.border = THIN_BORDER
    ws.row_dimensions[3].height = 25

    # 数据行
    for row_idx, row_data in enumerate(data_rows, 4):
        for col_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = DATA_FONT
            cell.border = THIN_BORDER
            if col_idx == 1 or (isinstance(val, (int, float)) and not isinstance(val, bool)):
                cell.alignment = CENTER_ALIGN
            else:
                cell.alignment = LEFT_ALIGN
            if row_idx % 2 == 0:
                cell.fill = ALT_FILL

    # 列宽
    if col_widths:
        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # 冻结窗格
    ws.freeze_panes = 'A4'
    # 自动筛选
    ws.auto_filter.ref = f'A3:{get_column_letter(ncols)}{3 + len(data_rows)}'

    return wb


def extract_region(address):
    """从地址中提取所在地区"""
    if not address:
        return '未知'
    address = str(address)
    # 匹配河北省内城市
    cities = ['石家庄', '唐山', '秦皇岛', '邯郸', '邢台', '保定', '张家口', '承德', '沧州', '廊坊', '衡水']
    for c in cities:
        if c in address:
            return c
    # 北京/天津
    if '北京' in address:
        return '北京'
    if '天津' in address:
        return '天津'
    # 省外
    return '省外'


def process_2024():
    # 创建输出目录
    subdirs = ['分数线', '招生计划', '一分一档', '院校详情', '专业计划及学费']
    for sd in subdirs:
        os.makedirs(os.path.join(OUT_DIR, sd), exist_ok=True)

    total_stats = {'分数线': 0, '招生计划': 0, '一分一档': 0, '院校详情': 0, '专业计划及学费': 0}

    for src_suffix, out_cat, full_name, cat_num in CATEGORIES:
        src_file = os.path.join(SRC_DIR, f'2024河北单招_{src_suffix}_数据.xlsx')
        if not os.path.exists(src_file):
            print(f'[跳过] 源文件不存在: {src_file}')
            continue

        print(f'[处理] 2024年 {out_cat} ...')
        wb_src = openpyxl.load_workbook(src_file, data_only=True)

        # ===== 1. 读取各sheet数据 =====

        # 分数线（只有控制线）
        score_line_data = []
        control_line = None
        if '分数线' in wb_src.sheetnames:
            ws = wb_src['分数线']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] == '控制线(分)':
                    control_line = row[1]
                    break

        # 招生计划汇总
        plan_summary = {}
        if '招生计划汇总' in wb_src.sheetnames:
            ws = wb_src['招生计划汇总']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0] and row[1]:
                    plan_summary[row[0]] = row[1]

        # 院校详情
        school_details = []
        if '院校详情' in wb_src.sheetnames:
            ws = wb_src['院校详情']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1]:  # 院校名称
                    school_details.append({
                        '序号': row[0],
                        '院校名称': row[1],
                        '办学性质': row[2] or '未知',
                        '院校地址': row[3] or '',
                    })

        # 一分一档
        yfyd_data = []
        if '一分一档' in wb_src.sheetnames:
            ws = wb_src['一分一档']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[0]:
                    yfyd_data.append({
                        '分数': row[0],
                        '本段人数': row[1],
                        '累计人数': row[2],
                        '累计占比': row[3],
                    })

        # 专业计划及学费
        major_plan_data = []
        if '专业计划及学费' in wb_src.sheetnames:
            ws = wb_src['专业计划及学费']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1]:  # 院校名称
                    major_plan_data.append({
                        '序号': row[0],
                        '院校名称': row[1],
                        '专业名称': row[2],
                        '考试类': row[3],
                        '计划数': row[4],
                        '学制': row[5],
                        '学费': row[6],
                        '备注': row[7],
                    })

        # 院校地址
        school_addr = {}
        if '院校地址' in wb_src.sheetnames:
            ws = wb_src['院校地址']
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[1]:
                    school_addr[row[1]] = row[2] or ''

        # ===== 2. 生成分数线 =====
        # 2024年只有省控制线，没有院校专业投档分，生成控制线说明文件
        score_rows = []
        if control_line:
            score_rows = [
                ('—', '—', f'考试{out_cat}', f'省控制线：{control_line}分（2024年未公布院校专业投档分）')
            ]
        wb_out = create_workbook(
            f'2024年河北高职单招录取分数线 - {full_name}',
            f'数据来源：河北省教育考试院 | 省控制线：{control_line}分 | 2024年未公布院校专业级投档分',
            ['院校名称', '专业名称', '考试类', '录取投档分'],
            score_rows,
            col_widths=[25, 35, 15, 40]
        )
        out_path = os.path.join(OUT_DIR, '分数线', f'2024年河北单招_{out_cat}_分数线.xlsx')
        wb_out.save(out_path)
        total_stats['分数线'] += 1

        # ===== 3. 生成招生计划 =====
        plan_rows = []
        for item in major_plan_data:
            plan_rows.append((
                item['院校名称'],
                item['专业名称'],
                item['考试类'],
                item['计划数'],
                item['学制'],
                item['学费'],
            ))
        total_plan = sum(int(x['计划数']) for x in major_plan_data if x['计划数'] and str(x['计划数']).isdigit())
        wb_out = create_workbook(
            f'2024年河北高职单招招生计划 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(plan_rows)} 条记录 | 总计划数: {total_plan} 人',
            ['院校名称', '专业名称', '考试类', '计划数', '学制', '学费(元/年)'],
            plan_rows,
            col_widths=[25, 35, 12, 10, 8, 12]
        )
        out_path = os.path.join(OUT_DIR, '招生计划', f'2024年河北单招_{out_cat}_招生计划.xlsx')
        wb_out.save(out_path)
        total_stats['招生计划'] += 1

        # ===== 4. 生成一分一档 =====
        yfyd_rows = []
        for item in yfyd_data:
            yfyd_rows.append((
                item['分数'],
                item['本段人数'],
                item['累计人数'],
                item['累计占比'],
            ))
        total_people = yfyd_data[-1]['累计人数'] if yfyd_data else 0
        wb_out = create_workbook(
            f'2024年河北高职单招一分一档表 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(yfyd_rows)} 条分数段 | 总人数: {total_people} 人',
            ['分数', '本段人数', '累计人数', '累计占比'],
            yfyd_rows,
            col_widths=[15, 12, 12, 12]
        )
        out_path = os.path.join(OUT_DIR, '一分一档', f'2024年河北单招_{out_cat}_一分一档.xlsx')
        wb_out.save(out_path)
        total_stats['一分一档'] += 1

        # ===== 5. 生成院校详情 =====
        school_rows = []
        for idx, item in enumerate(school_details, 1):
            region = extract_region(item.get('院校地址', ''))
            nature = item.get('办学性质', '未知')
            # 规范化办学性质
            if '公办' in str(nature):
                nature = '公办'
            elif '民办' in str(nature):
                nature = '民办'
            school_rows.append((
                idx,
                item['院校名称'],
                region,
                nature,
                '招生计划 院校资讯',
            ))
        wb_out = create_workbook(
            f'2024年河北高职单招招生院校详情 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(school_rows)} 所院校',
            ['序号', '院校名称', '所在地区', '办学性质', '操作'],
            school_rows,
            col_widths=[8, 30, 12, 12, 20]
        )
        out_path = os.path.join(OUT_DIR, '院校详情', f'2024年河北单招_{out_cat}_院校详情.xlsx')
        wb_out.save(out_path)
        total_stats['院校详情'] += 1

        # ===== 6. 生成专业计划及学费（增强版，含地址） =====
        major_rows = []
        for idx, item in enumerate(major_plan_data, 1):
            addr = school_addr.get(item['院校名称'], '')
            major_rows.append((
                idx,
                item['院校名称'],
                addr,
                item['专业名称'],
                item['计划数'],
                item['学制'],
                item['学费'],
                '—',  # 2024年无投档分
            ))
        wb_out = create_workbook(
            f'2024年河北高职单招专业计划及学费 - {full_name}',
            f'数据来源：河北省教育考试院 | 共 {len(major_rows)} 个专业 | 总计划数: {total_plan} 人 | 2024年未公布投档分',
            ['序号', '院校名称', '院校地址', '专业名称', '计划数', '学制', '学费(元/年)', '2024投档最低分'],
            major_rows,
            col_widths=[8, 22, 30, 30, 10, 8, 12, 14]
        )
        out_path = os.path.join(OUT_DIR, '专业计划及学费', f'2024年河北单招_{out_cat}_专业计划及学费.xlsx')
        wb_out.save(out_path)
        total_stats['专业计划及学费'] += 1

        print(f'  ✓ 分数线1 招生计划{len(plan_rows)} 一分一档{len(yfyd_rows)} 院校{len(school_rows)} 专业{len(major_rows)}')

    print(f'\n===== 2024年数据转换完成 =====')
    for k, v in total_stats.items():
        print(f'  {k}: {v}个文件')


if __name__ == '__main__':
    process_2024()
