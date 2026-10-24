# -*- coding: utf-8 -*-
"""
将2026年单招Excel数据转换为JS格式，供网页使用
"""
import openpyxl
import json
import os

BASE = r'C:\Users\XIAO\Desktop\大类信息\2026单招数据'
OUT_JS = os.path.join(r'C:\Users\XIAO\Desktop\大类信息\单招数据网页', 'data.js')

CATEGORIES = [
    ('1', '一类', '考试一类', '土木建筑、资源环境与安全、水利等'),
    ('2', '二类', '考试二类', '交通运输、能源动力与材料等'),
    ('3', '三类', '考试三类', '装备制造等'),
    ('4', '四类', '考试四类', '农林牧渔、轻工纺织、生物与化工、食品药品与粮食等'),
    ('5', '五类', '考试五类', '公共管理与服务、旅游、公安与司法等'),
    ('6', '六类', '考试六类', '教育与体育、新闻传播等'),
    ('7', '七类', '考试七类', '医药卫生（临床医学类专业除外）等'),
    ('8', '八类美术', '考试八类-美术设计类', '文化艺术等'),
    ('9', '九类', '考试九类', '财经商贸等'),
    ('10', '十类', '考试十类', '电子信息、计算机等'),
]


def read_excel_data(filepath, header_row=3):
    """读取Excel数据，返回表头和数据行"""
    if not os.path.exists(filepath):
        return [], []
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active
    headers = []
    for col in range(1, ws.max_column + 1):
        v = ws.cell(row=header_row, column=col).value
        if v:
            headers.append(str(v).strip())
    data = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if all(v is None for v in row):
            continue
        row_data = {}
        for i, h in enumerate(headers):
            val = row[i] if i < len(row) else None
            if val is not None:
                row_data[h] = str(val).strip() if isinstance(val, str) else val
            else:
                row_data[h] = ''
        data.append(row_data)
    return headers, data


def main():
    os.makedirs(os.path.dirname(OUT_JS), exist_ok=True)
    result = {}

    for cat_id, cat_name, cat_full, cat_desc in CATEGORIES:
        print(f'[转换] {cat_name} ...')
        cat_data = {
            'id': cat_id,
            'name': cat_name,
            'fullName': cat_full,
            'desc': cat_desc,
        }

        # 分数线
        _, score_data = read_excel_data(os.path.join(BASE, '分数线', f'2026年河北单招_{cat_name}_分数线.xlsx'))
        cat_data['scoreLines'] = score_data

        # 招生计划
        _, plan_data = read_excel_data(os.path.join(BASE, '招生计划', f'2026年河北单招_{cat_name}_招生计划.xlsx'))
        cat_data['plans'] = plan_data

        # 一分一档
        _, yfyd_data = read_excel_data(os.path.join(BASE, '一分一档', f'2026年河北单招_{cat_name}_一分一档.xlsx'))
        cat_data['yifenyidang'] = yfyd_data

        # 院校详情
        _, school_data = read_excel_data(os.path.join(BASE, '院校详情', f'2026年河北单招_{cat_name}_院校详情.xlsx'))
        cat_data['schools'] = school_data

        # 专业计划及学费
        _, major_data = read_excel_data(os.path.join(BASE, '专业计划及学费', f'2026年河北单招_{cat_name}_专业计划及学费.xlsx'))
        cat_data['majors'] = major_data

        result[cat_id] = cat_data

        print(f'  分数线{len(score_data)} 招生计划{len(plan_data)} 一分一档{len(yfyd_data)} 院校{len(school_data)} 专业{len(major_data)}')

    # 写入JS
    js_content = 'const DANZHAO_DATA_2026 = ' + json.dumps(result, ensure_ascii=False, separators=(',', ':')) + ';\n'

    with open(OUT_JS, 'w', encoding='utf-8') as f:
        f.write(js_content)

    size_kb = os.path.getsize(OUT_JS) / 1024
    print(f'\n===== 转换完成 =====')
    print(f'输出: {OUT_JS}')
    print(f'大小: {size_kb:.1f} KB')


if __name__ == '__main__':
    main()
